# -*- coding: utf-8 -*-
"""사용 설명서 그림 만들기 — 예제 자료로 도우미를 띄워 화면을 찍고 설명 번호를 그려 넣는다.

    python tools/manual_shots.py [--forms-origin http://127.0.0.1:8787]      → docs/manual_img/*.png

준비
  · Edge 또는 Chrome 이 설치된 Windows, .venv 에 playwright (pip install -r requirements-docs.txt)
  · 휴대전화 입력 페이지 그림과 데이터 입력 관리 그림은 오토다타 웹 서버가 필요하다 — 저장소 루트에서
    `npx wrangler@4 dev --config wrangler.jsonc --port 8787` 를 켜 두면 실서버 없이 만든다(끝나면 만든 양식을 지운다).
실제 data/ 는 건드리지 않는다: 임시 작업 폴더에 예제 템플릿·양식·기록을 만든다. 그림의 번호는 docs/manual.md 의
{{1}}, {{2}} … 설명과 순서가 같아야 한다.
"""
from __future__ import annotations

import argparse
import base64
import http.server
import io
import json
import os
import shutil
import socket
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from functools import partial
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "docs" / "manual_img"
EXAMPLES = ROOT / "examples"
WORK = Path(tempfile.mkdtemp(prefix="autodata_manual_"))
CALLOUT = (194, 37, 92)
WIDTH = 1180                       # 작업 화면 그림 너비
PRETTY_ORIGIN = "https://autodata.singlena.workers.dev"
TEMPLATE = "야생조류 현장 조사표"
FORM_TITLE = "가온천 봄철 탐조 (예제)"


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


def free_port(start: int) -> int:
    for p in range(start, start + 30):
        with socket.socket() as s:
            s.settimeout(0.2)
            if s.connect_ex(("127.0.0.1", p)) != 0:
                return p
    raise SystemExit("빈 포트가 없습니다")


def http_json(method, url, body=None, headers=None, timeout=60):
    data = None
    h = {"User-Agent": "AutoData-manual", **(headers or {})}
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        h["content-type"] = "application/json"
    req = urllib.request.Request(url, data=data, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()
    except (urllib.error.URLError, OSError) as e:          # 아직 안 뜬 서버 등 — 상태 0
        return 0, str(e).encode("utf-8")


# ---------------------------------------------------------------- 도우미(임시 자료로) 띄우기
def start_helper(forms_origin: str) -> int:
    os.environ["AUTODATA_FORMS_ORIGIN"] = forms_origin
    from app import config
    data = WORK / "data"
    for name in ("uploads", "output", "pdf_cache", "report_cache", "template_pdfs"):
        (data / name).mkdir(parents=True, exist_ok=True)
    config.BASE_DIR = WORK                      # ai_config.txt·forms_config.txt 등 실제 설정을 읽지 않게
    config.DATA_DIR = data
    config.UPLOAD_DIR, config.OUTPUT_DIR = data / "uploads", data / "output"
    config.PDF_CACHE_DIR, config.REPORT_CACHE_DIR = data / "pdf_cache", data / "report_cache"
    config.TEMPLATE_PDF_DIR = data / "template_pdfs"
    config.SETTINGS_PATH, config.SCE_CONFIG_PATH = WORK / "ai_settings.enc", WORK / "sce_config.json"
    for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GEMINI_API_KEY", "FIELD_SURVEY_PROXY_URL", "FIELD_SURVEY_APP_TOKEN"):
        os.environ.pop(k, None)
    config.reload_ai_settings()
    import app.main as app_main
    from core import forms
    from core.accumulate import AccumulateStore
    from core.review.store import CorrectionStore
    from core.template.store import TemplateStore
    app_main._TEMPLATES = TemplateStore(data / "templates.json")
    app_main._FORMS = forms.FormRegistry(data / "forms.json", data / "forms")
    app_main._STORE = CorrectionStore(data / "corrections.json")
    app_main._ACCUM = AccumulateStore(data / "output" / "누적DB.json")
    # 예제 템플릿(양식 PDF 포함)
    boxes = json.loads((EXAMPLES / "template_boxes.json").read_text(encoding="utf-8"))
    app_main._TEMPLATES.save(TEMPLATE, boxes)
    shutil.copyfile(EXAMPLES / "야생조류_현장조사표_빈양식.pdf", app_main._tpl_pdf_path(TEMPLATE))
    import uvicorn
    port = free_port(8765)
    server = uvicorn.Server(uvicorn.Config(app_main.app, host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    for _ in range(100):
        time.sleep(0.2)
        st, _ = http_json("GET", f"http://127.0.0.1:{port}/health", timeout=3)
        if st == 200:
            return port
    raise SystemExit("도우미가 뜨지 않았습니다")


def start_web_page() -> str | None:
    """웹 시작 페이지(web/)를 8788 에서 — 도우미가 이 주소의 상태 읽기를 허용한다."""
    handler = partial(http.server.SimpleHTTPRequestHandler, directory=str(ROOT / "web"))
    handler.log_message = lambda *a, **k: None
    try:
        httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 8788), handler)
    except OSError:
        return None
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return "http://localhost:8788"


# ---------------------------------------------------------------- 그림 저장(번호 표시)
def _font(size: int):
    from PIL import ImageFont
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    for name in ("malgunbd.ttf", "arialbd.ttf"):
        try:
            return ImageFont.truetype(str(fonts / name), size)
        except OSError:
            continue
    return ImageFont.load_default()


class Shooter:
    def __init__(self, page, width: int = WIDTH):
        self.page = page
        self.width = width
        self.scale = float(page.evaluate("window.devicePixelRatio") or 1)   # 그림 픽셀 = CSS 픽셀 × 배율
        self.font = _font(int(17 * self.scale))
        self.r = int(15 * self.scale)
        OUT.mkdir(parents=True, exist_ok=True)

    def fit(self, cap: int = 3200) -> None:
        """스크롤 없이 문서 전체가 보이게 창 높이를 맞춘다(좌표 = 문서 좌표)."""
        js = """() => Math.ceil(Math.max(0, ...[...document.body.children]
                   .filter(e => e.tagName !== 'SCRIPT' && !e.hidden && getComputedStyle(e).display !== 'none'
                                && getComputedStyle(e).position !== 'fixed')
                   .map(e => e.getBoundingClientRect().bottom + window.scrollY))) + 28"""
        last = -1
        for _ in range(6):        # 창 높이를 바꾸면 늘어나는 칸(flex)이 있어 높이가 안정될 때까지 반복
            h = int(min(max(self.page.evaluate(js), 500), cap))
            if abs(h - last) <= 2:
                break
            self.page.set_viewport_size({"width": self.width, "height": h})
            self.page.wait_for_timeout(250)
            last = h

    def rect(self, selector, pad: int = 4, nth: int = 0):
        loc = self.page.locator(selector)
        if loc.count() <= nth:
            log(f"경고: '{selector}' 를 찾지 못했습니다")
            return None
        b = loc.nth(nth).bounding_box()
        if not b:
            return None
        k = self.scale
        return ((b["x"] - pad) * k, (b["y"] - pad) * k, (b["x"] + b["width"] + pad) * k, (b["y"] + b["height"] + pad) * k)

    @staticmethod
    def union(*rects):
        rects = [r for r in rects if r]
        return (min(r[0] for r in rects), min(r[1] for r in rects), max(r[2] for r in rects), max(r[3] for r in rects))

    def grab(self):
        from PIL import Image
        return Image.open(io.BytesIO(self.page.screenshot())).convert("RGB")

    def shot(self, name: str, marks=(), crop=None, img=None, start: int = 1, save: bool = True):
        from PIL import ImageDraw
        img = img or self.grab()
        d = ImageDraw.Draw(img)
        W, H = img.size
        for i, m in enumerate(marks, start):
            if m is None:         # 못 찾은 대상 — 번호는 건너뛰지 않는다(설명 번호와 맞추기 위해)
                log(f"경고: {name} 의 {i}번 표시 대상을 찾지 못했습니다 — 그림에 번호가 빠집니다.")
                continue
            l, t, r, b = m
            l, t, r, b = max(l, 3), max(t, 3), min(r, W - 4), min(b, H - 4)
            if r - l < 8 or b - t < 8:
                log(f"경고: {name} 의 {i}번 표시 대상이 화면 밖에 있습니다 — 그림과 설명 번호가 어긋납니다.")
                continue
            rr = self.r
            d.rounded_rectangle((l, t, r, b), radius=6, outline=CALLOUT, width=int(3 * self.scale))
            cx, cy = min(max(l, rr + 2), W - rr - 2), min(max(t, rr + 2), H - rr - 2)
            d.ellipse((cx - rr, cy - rr, cx + rr, cy + rr), fill=CALLOUT, outline="white", width=2)
            d.text((cx, cy), str(i), fill="white", font=self.font, anchor="mm")
        if crop:
            img = img.crop((int(max(0, crop[0])), int(max(0, crop[1])), int(min(W, crop[2])), int(min(H, crop[3]))))
        if save:
            img.save(OUT / f"{name}.png", optimize=True)
            log("그림:", f"{name}.png", img.size)
        return img


# ---------------------------------------------------------------- 각 화면
def shot_web_start(s: Shooter, web: str | None, port: int):
    if not web:
        log("웹 시작 페이지(8788)를 열 수 없어 web_start.png 를 건너뜁니다")
        return
    p = s.page
    p.goto(f"{web}/?port={port}")
    try:
        p.wait_for_selector('li[data-key="helper"][data-state="ok"]', timeout=15000)
    except Exception:  # noqa: BLE001
        log("경고: 시작 페이지가 도우미를 찾지 못한 채로 찍습니다")
    p.wait_for_timeout(800)
    s.fit()
    s.shot("web_start", [s.rect(".checks"), s.rect("#open"), s.rect("#install"), s.rect("#manual")])


def shot_start(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/")
    p.wait_for_selector("a.card")
    s.fit()
    marks = [s.rect(".tabs")] + [s.rect("a.card", nth=i) for i in range(6)] + [s.rect("#manualLink")]
    s.shot("start", marks)


def shot_designer(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/pdf-designer")
    p.wait_for_selector("#tplStart:not([hidden])", timeout=15000)
    s.fit()
    s.shot("designer_load", [s.rect("#dropzone"), s.rect("#tplStart")], crop=(0, 0, WIDTH, s.rect("#loadCard", 16)[3]))
    p.select_option("#startTplSel", TEMPLATE)
    p.click("#startTplBtn")
    p.wait_for_selector("#main:not([hidden])", timeout=20000)
    p.wait_for_function("document.querySelectorAll('.pbox').length >= 10", timeout=30000)
    p.wait_for_function("[...document.querySelectorAll('#pageHost img')].every(i => i.complete && i.naturalWidth > 0)", timeout=30000)
    p.wait_for_timeout(600)
    s.fit()
    canvas = s.rect("#pageHost")
    marks = [s.rect(".tools-row"), s.rect("#pageNav"),
             (canvas[0], canvas[1], canvas[2], min(canvas[3], canvas[1] + 620)),
             s.union(s.rect(".box-pane .pane-head"), s.rect("#boxList")), s.rect(".save-row"),
             s.rect(".steps-bottom .step-card", nth=0), s.rect(".next-step")]
    top = s.rect(".grid-pane .pane-head", 24)
    s.shot("designer", marks, crop=(0, top[1], WIDTH, 99999))
    # 여러 박스 선택 → 선택 막대
    boxes = p.locator(".pbox")
    boxes.nth(1).click()
    boxes.nth(4).click(modifiers=["Shift"])
    p.wait_for_selector("#selBar:not([hidden])", timeout=5000)
    p.wait_for_timeout(300)
    s.fit()
    pane = s.rect(".box-pane", 8)
    crop = (pane[0], pane[1], pane[2], min(pane[3], pane[1] + 560))
    s.shot("designer_selbar", [s.rect("#selBar"), s.rect("#boxList li.sel", nth=0)], crop=crop)


def _entries():
    return json.loads((EXAMPLES / "entries.json").read_text(encoding="utf-8"))


def _photo_data_url() -> str:
    return "data:image/jpeg;base64," + base64.b64encode((EXAMPLES / "현장사진_예제.jpg").read_bytes()).decode("ascii")


def phone_submit(share_url: str, entry: dict, photo: bool):
    origin, rest = share_url.split("/f/")
    fid, key = rest.split("?k=")
    body = {"id": entry["id"], "client_id": entry["client_id"], "created": entry["created"], "values": entry["values"],
            "photos": {"현장사진": _photo_data_url()} if photo else {}, "meta": entry.get("meta") or {}}
    st, payload = http_json("POST", f"{origin}/api/forms/{fid}/entries?k={key}", body)
    if st not in (200, 201):
        raise SystemExit(f"기록 보내기 실패 {st} {payload[:200]!r}")


def shot_entry(s: Shooter, base: str) -> dict | None:
    p = s.page
    p.goto(base + "/entry")
    p.wait_for_selector("#tplSel option", state="attached")
    s.fit()
    card = s.rect(".steps-bottom .step-card", 16, nth=0)
    s.shot("entry_publish", [s.rect("#tplSel"), s.rect("#titleInput"), s.rect("#gpsChk", 6), s.rect("#publishBtn")],
           crop=(0, card[1], WIDTH, card[3]))
    p.select_option("#tplSel", TEMPLATE)
    p.fill("#titleInput", FORM_TITLE)
    p.check("#gpsChk")
    p.click("#publishBtn")
    try:
        p.wait_for_selector("#formDetail:not([hidden])", timeout=40000)
    except Exception:  # noqa: BLE001
        msg = p.locator("#publishMsg").inner_text()
        log("경고: 양식을 만들지 못해 데이터 입력 관리 그림을 건너뜁니다 —", msg)
        return None
    form = p.evaluate("FORMS[0]")
    share = form["share_url"]
    # 현장(휴대전화)에서 보낸 기록 — 예제 3건 중 2건은 여기서 미리 보내고, 1건은 휴대전화 화면에서 보낸다
    ents = _entries()
    phone_submit(share, ents[0], photo=True)
    phone_submit(share, ents[1], photo=False)
    return {"form": form, "share": share, "third": ents[2]}


def shot_phone(pw, share: str, entry: dict):
    """휴대전화 입력 페이지 — 위쪽(항목)과 아래쪽(표·위치·보낸 기록·버튼)을 나란히 붙인 한 장."""
    from PIL import Image
    ctx = pw.chromium.launch(channel=CHANNEL, headless=True).new_context(
        viewport={"width": 390, "height": 1040}, device_scale_factor=2, is_mobile=True, has_touch=True, locale="ko-KR",
        bypass_csp=True)   # 입력 페이지의 CSP 가 Playwright 의 평가 스크립트를 막지 않게
    p = ctx.new_page()
    p.goto(share)
    p.wait_for_selector("#sendBtn")
    v = entry["values"]
    # 먼저 한 건을 보내(보낸 기록 목록에 남게) — 실제로 웹에 저장된다
    p.fill('[data-k="조사지역"]', v["조사지역"])
    p.fill('[data-k="조사일자"]', v["조사일자"])
    p.fill('[data-k="조사자"]', v["조사자"])
    p.fill('[data-k="날씨"]', v["날씨"])
    p.fill('[data-k="시작시각"]', v["시작시각"])
    p.fill('[data-k="종료시각"]', v["종료시각"])
    for k in ("도보", "차량", "선박"):
        if v.get(k):
            p.check(f'[data-k="{k}"]')
    rows = v["관찰기록"]
    for i, r in enumerate(rows):
        if i > 0:
            p.click('[data-addrow="관찰기록"]')
        tr = p.locator('[data-table="관찰기록"] tbody tr').nth(i)
        for col, val in r.items():
            tr.locator(f'input[data-col="{col}"]').fill(str(val))
    p.fill('[data-k="특이사항"]', v["특이사항"])
    p.fill('[data-k="사진설명"]', v["사진설명"])
    p.fill('[data-k="조사기관"]', v["조사기관"])
    p.fill('[data-k="확인자"]', v["확인자"])
    p.evaluate("""() => { const o = document.getElementById('gpsOut'); if (o) o.textContent = '37.0988, 127.5012 (±12m)'; }""")
    p.click("#sendBtn")
    p.wait_for_function("document.getElementById('status').textContent.includes('보냈습니다')", timeout=20000)
    # 다음 기록을 쓰는 중인 화면 — 위쪽
    p.fill('[data-k="조사지역"]', "가온천 하류 습지")
    p.fill('[data-k="조사일자"]', "2026-04-15")
    p.fill('[data-k="조사자"]', "홍길동")
    p.fill('[data-k="날씨"]', "맑음")
    p.check('[data-k="도보"]')
    p.evaluate("window.scrollTo(0, 0)")
    p.wait_for_timeout(300)
    s = Shooter(p)
    top = s.grab()
    marks_top = [s.rect(".band"), s.union(s.rect('[data-k="조사지역"]', 2), s.rect('[data-k="시작시각"]', 2)),
                 s.rect(".chk", nth=0)]
    top = s.shot("phone_top", marks_top, img=top, save=False)
    # 아래쪽: 표·사진·위치·보낸 기록·버튼 — 표(여러 행)가 맨 위에 오도록 스크롤
    p.locator('[data-table="관찰기록"]').evaluate("el => el.scrollIntoView({block: 'start'})")
    p.evaluate("window.scrollBy(0, -12)")
    p.wait_for_timeout(300)
    bottom = s.grab()
    marks_bottom = [s.union(s.rect('[data-table="관찰기록"] table', 2), s.rect('[data-addrow="관찰기록"]', 2)),
                    s.rect(".pbtn", nth=0), s.rect("#gpsBtn"), s.rect("#hist"), s.rect(".bar")]
    bottom = s.shot("phone_bottom", marks_bottom, img=bottom, start=4, save=False)
    gap = 24
    canvas = Image.new("RGB", (top.width + bottom.width + gap, max(top.height, bottom.height)), (238, 242, 236))
    canvas.paste(top, (0, 0))
    canvas.paste(bottom, (top.width + gap, 0))
    canvas.save(OUT / "phone.png", optimize=True)
    log("그림: phone.png", canvas.size)
    ctx.close()


def shot_entry_live(s: Shooter, base: str, form: dict):
    p = s.page
    p.goto(base + "/entry")
    p.wait_for_selector("#formList .tpl-load")
    p.click("#formList .tpl-load")
    p.wait_for_selector("#formDetail:not([hidden])")
    p.wait_for_function("document.getElementById('countPill').textContent.startsWith('3')", timeout=30000)
    p.wait_for_function("[...document.querySelectorAll('.thumb')].every(i => i.complete)", timeout=15000)
    # 그림에는 실서버 주소로(로컬 시험 서버 주소 대신) — QR 도 그 주소로 다시 만든다
    pretty = f"{PRETTY_ORIGIN}/f/{form['id']}?k={form['write_key']}"
    from core.forms import qr_svg
    qr = "data:image/svg+xml;base64," + base64.b64encode(qr_svg(pretty).encode("utf-8")).decode("ascii")
    p.evaluate("""([u, q]) => { document.getElementById('shareLink').value = u; document.getElementById('openLink').href = u;
                                document.getElementById('qrImg').src = q; }""", [pretty, qr])
    p.wait_for_timeout(500)
    s.fit()
    p.evaluate("() => { const el = document.querySelector('.entry-scroll'); el.scrollLeft = el.scrollWidth; }")
    p.wait_for_timeout(200)
    card2 = s.rect(".steps-bottom .step-card", 16, nth=1)
    s.shot("entry_share", [s.rect("#formList"), s.rect("#qrImg"), s.rect(".share-link-row"), s.rect(".share-actions")],
           crop=(0, card2[1], WIDTH, card2[3]))
    live = s.union(s.rect("#entriesCard", 16), s.rect("#exportCard", 16))
    s.shot("entry_live", [s.rect(".sync-row"), s.rect("#entryTable"), s.rect(".thumb", 3, nth=0), s.rect(".export-row")],
           crop=(0, live[1], WIDTH, live[3]))


def shot_filled_pdf(base: str, form: dict):
    """현장조사표 PDF 로 내보낸 첫 기록 — 도우미가 만든 PDF 의 첫 쪽."""
    import fitz
    st, pdf = http_json("GET", f"{base}/api/forms/{form['id']}/pdf?entry={_entries()[0]['id']}")
    if st != 200:
        log("경고: 조사표 PDF 를 만들지 못했습니다", st, pdf[:120])
        return
    doc = fitz.open(stream=pdf, filetype="pdf")
    dpi = 110
    pix = doc[0].get_pixmap(dpi=dpi)
    from PIL import Image, ImageDraw
    img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    k = dpi / 72
    boxes = {b["field"]: b for b in json.loads((EXAMPLES / "template_boxes.json").read_text(encoding="utf-8"))}

    def R(f, pad=3):
        b = boxes[f]
        return (b["x0"] * k - pad, b["y0"] * k - pad, b["x1"] * k + pad, b["y1"] * k + pad)
    marks = [R("조사지역"), R("도보"), R("관찰기록"), R("현장사진"), (22 * k, (842 - 18) * k, 380 * k, (842 - 2) * k)]
    d = ImageDraw.Draw(img)
    font = _font(17)
    for i, (l, t, r, b) in enumerate(marks, 1):
        d.rounded_rectangle((l, t, r, b), radius=6, outline=CALLOUT, width=3)
        d.ellipse((l - 15, t - 15, l + 15, t + 15), fill=CALLOUT, outline="white", width=2)
        d.text((l, t), str(i), fill="white", font=font, anchor="mm")
    img = img.crop((int(18 * k), int(25 * k), int(560 * k), int(842 * k)))
    img.save(OUT / "filled_pdf.png", optimize=True)
    log("그림: filled_pdf.png", img.size)


def shot_extract(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/extract")
    p.wait_for_selector("#tplUsed li")
    files = [str(EXAMPLES / f"작성예시_{i:02d}.pdf") for i in (1, 2, 3)]
    p.set_input_files("#applyInput", files)
    p.wait_for_timeout(300)
    s.fit()
    top = s.union(s.rect(".steps-bottom .step-card", 16, nth=0), s.rect(".steps-bottom .step-card", 16, nth=1))
    s.shot("extract", [s.rect("#tplUsed"), s.rect(".sheetname-row"), s.rect("#applyInput"), s.rect("#applyBtn")],
           crop=(0, top[1], WIDTH, top[3]))
    p.click("#applyBtn")
    p.wait_for_selector("#applyResult table.apply-table", timeout=180000)
    p.wait_for_timeout(500)
    s.fit()
    res = s.rect("#applyResult", 10)
    marks = [s.rect("#applyResult > p.muted", 2, nth=0), s.rect("#applyResult .btn-download"),
             s.rect("#applyResult table.apply-table", 2, nth=0)]
    extra = s.rect("#applyResult p.muted", 2, nth=1)
    if extra:
        marks.append(extra)
    s.shot("extract_result", marks, crop=(0, res[1], WIDTH, min(res[3], res[1] + 760)))
    # 보고서 출력 양식 — 자리표시자를 넣은 엑셀 양식을 올린 편집 화면
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "조사 결과"
    ws["A1"] = "야생조류 조사 결과 보고"
    for r, (k, v) in enumerate([("조사지역", "{조사지역}"), ("조사일자", "{조사일자}"), ("조사자", "{조사자}"),
                                ("날씨", "{날씨}"), ("특이사항", "{특이사항}")], start=3):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    ws["A9"], ws["B9"], ws["C9"] = "종명", "개체수", "행동"
    ws["A10"], ws["B10"], ws["C10"] = "{종명}", "{개체수}", "{행동}"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    tpl = WORK / "보고서양식_예제.xlsx"
    wb.save(tpl)
    p.set_input_files("#reportInput", str(tpl))
    p.wait_for_selector("#reportEditor:not([hidden])", timeout=30000)
    p.wait_for_selector("#rptGrid td", timeout=30000)
    p.wait_for_timeout(400)
    s.fit()
    card3 = s.rect(".steps-bottom .step-card", 16, nth=2)
    s.shot("report_form", [s.rect(".rpt-draft-row"), s.rect("#reportInput"), s.rect(".rpt-toolbar"), s.rect("#rptGrid"),
                           s.rect("#rptRunRow")], crop=(0, card3[1], WIDTH, min(card3[3], card3[1] + 900)))


def shot_ai(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/ai")
    p.wait_for_function("!document.getElementById('status').textContent.includes('확인 중')", timeout=15000)
    p.wait_for_timeout(500)
    s.fit()
    s.shot("ai", [s.rect("#status"), s.union(s.rect("#files"), s.rect("#run")), s.rect("#accumCard")])


def shot_settings(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/settings")
    p.wait_for_function("document.getElementById('active').textContent.length > 0", timeout=15000)
    p.wait_for_function("!document.getElementById('rt_body').textContent.includes('확인 중')", timeout=20000)
    p.wait_for_timeout(500)
    s.fit()
    s.shot("settings", [s.rect(".card", nth=0), s.rect(".card", nth=1), s.rect(".card", nth=2), s.rect(".card", nth=3),
                        s.rect("#ocr")])


def shot_system(s: Shooter, base: str):
    p = s.page
    p.goto(base + "/system")
    p.wait_for_function("document.querySelectorAll('#rows tr').length > 3", timeout=60000)
    p.wait_for_function("!document.getElementById('overall').textContent.includes('점검하는 중')", timeout=60000)
    p.wait_for_timeout(500)
    # 그림에 임시 작업 폴더 경로가 보이지 않게(표시만 바꾼다)
    p.evaluate("""(work) => { for (const td of document.querySelectorAll('#rows td'))
                     if (td.textContent.includes(work)) td.textContent = td.textContent.split(work).join('C:\AutoData'); }""", str(WORK))
    s.fit()
    s.shot("system", [s.rect("#overall"), s.rect("#tbl"), s.rect(".btns"), s.rect(".card", nth=1)])


CHANNEL = "msedge"


def main():
    global CHANNEL
    ap = argparse.ArgumentParser(description="사용 설명서 그림 만들기")
    ap.add_argument("--forms-origin", default="http://127.0.0.1:8787", help="디지털 입력 양식 서버(wrangler dev)")
    ap.add_argument("--browser", default="msedge", choices=["msedge", "chrome"])
    a = ap.parse_args()
    CHANNEL = a.browser
    if not (EXAMPLES / "template_boxes.json").exists():
        from tools.make_examples import main as make
        make(EXAMPLES)
    st, _ = http_json("GET", a.forms_origin + "/api/forms/abcdefghij?k=x", timeout=5)
    if st != 404:
        log(f"경고: 양식 서버({a.forms_origin})가 없어 데이터 입력 관리·휴대전화 그림은 건너뜁니다 — "
            "저장소 루트에서 `npx wrangler@4 dev --config wrangler.jsonc --port 8787` 를 켜 두세요")
        forms_ok = False
    else:
        forms_ok = True
    port = start_helper(a.forms_origin)
    base = f"http://127.0.0.1:{port}"
    web = start_web_page()
    log("도우미", base, "· 웹 시작 페이지", web, "· 작업 폴더", WORK)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel=CHANNEL, headless=True)
        ctx = browser.new_context(viewport={"width": WIDTH, "height": 900}, device_scale_factor=1, locale="ko-KR")
        page = ctx.new_page()
        s = Shooter(page)
        shot_web_start(s, web, port)
        shot_start(s, base)
        shot_designer(s, base)
        made = shot_entry(s, base) if forms_ok else None
        if made:
            shot_phone(pw, made["share"], made["third"])
            shot_entry_live(s, base, made["form"])
            shot_filled_pdf(base, made["form"])
        shot_extract(s, base)
        shot_ai(s, base)
        shot_settings(s, base)
        shot_system(s, base)
        if made:   # 시험 서버의 양식은 지운다
            http_json("POST", f"{base}/api/forms/{made['form']['id']}/delete")
        browser.close()
    log("완료:", OUT)


if __name__ == "__main__":
    main()
