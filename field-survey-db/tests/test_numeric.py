"""'숫자' 유형 — 값 규칙과 엑셀 기록(수로 저장), 수치형 항목 자동 제안.

숫자로 지정한 열은 엑셀에 수(int/float)로 들어가야 합계·평균·정렬이 된다.
지정하지 않은 열은 문자 그대로 둔다('001' 같은 번호가 1이 되지 않게).
"""
import fitz
import openpyxl
import pytest

from core.numeric import as_number, excel_cell_value, first_number_text, looks_numeric_field


@pytest.mark.parametrize("raw,want", [
    ("약 25.5 m", "25.5"), ("30", "30"), ("1,234", "1234"),
    ("５０", "50"),                     # 전각 숫자(한글 서식·스캔본)
    (".5", "0.5"), ("-1.2", "-1.2"), ("+3", "3"),
    ("12개소", "12"), ("12.5%", "12.5"),
    ("없음", ""), ("", ""), (None, ""),
])
def test_first_number_text(raw, want):
    assert first_number_text(raw) == want


@pytest.mark.parametrize("raw,want", [
    ("30", 30), ("25.5", 25.5), ("1,234", 1234), ("５０", 50), ("0.5", 0.5),
    ("-1.2", -1.2), ("+3", 3),
    ("", None), ("없음", None),
    ("좌안 12 우안 15", None),          # 일부만 숫자면 문자로 둔다(값이 잘리지 않게)
    ("25.5 m", None),
])
def test_as_number(raw, want):
    got = as_number(raw)
    assert got == want
    if want is not None:
        assert isinstance(got, (int, float))


def test_as_number_keeps_int_and_float_apart():
    assert isinstance(as_number("30"), int) and isinstance(as_number("30.0"), float)


def test_excel_cell_value_only_for_number_columns():
    assert excel_cell_value("001", numeric=False) == "001"      # 일반 열은 그대로
    assert excel_cell_value("001", numeric=True) == 1
    assert excel_cell_value("없음", numeric=True) == "없음"      # 숫자 아니면 글자 유지


@pytest.mark.parametrize("name,want", [
    ("보 길이", True), ("길이 (m)", True), ("높이 (m)", True), ("유폭 m", True),
    ("월류수심", True), ("강수량", True), ("개체수", True), ("피도(%)", True),
    ("구조물번호", False), ("조사일시", False), ("위도", False), ("연락처", False),
    ("어도유무", False), ("재질_돌", False), ("비고", False), ("하천명", False),
])
def test_looks_numeric_field(name, want):
    assert looks_numeric_field(name) is want


def _form(path, rows, y_top=60):
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    font = fitz.Font("cjk")
    tw = fitz.TextWriter(page.rect)
    tw.append((60, 34), "인공구조물 현장 조사표", font=font, fontsize=18)
    x0, x1, x2, rh = 40, 160, 360, 28
    for i, (label, value) in enumerate(rows):
        y = y_top + i * rh
        page.draw_rect(fitz.Rect(x0, y, x1, y + rh), color=(0, 0, 0), width=0.8)
        page.draw_rect(fitz.Rect(x1, y, x2, y + rh), color=(0, 0, 0), width=0.8)
        tw.append((x0 + 6, y + 18), label, font=font, fontsize=10)
        if value:
            tw.append((x1 + 6, y + 18), value, font=font, fontsize=10)
    tw.write_text(page)
    doc.save(str(path))
    doc.close()


ROWS = [("하천명", "해남천"), ("보 길이 (m)", "25.5 m"), ("배사구 높이", "약 1,200"),
        ("구조물번호", "001"), ("어도유무", "있음")]


def test_suggest_marks_numeric_fields(tmp_path):
    """자동 제안 — 이름이 수치형인 값 칸은 '숫자' 유형으로, 나머지는 '일반'으로."""
    from core.pdf_pipeline import suggest_cells_maximal

    p = tmp_path / "form.pdf"
    _form(p, ROWS)
    by_field = {}
    for b in suggest_cells_maximal(str(p), 0):
        if (b.get("anchor") or {}).get("label"):        # 값 칸만(라벨 칸 제외)
            by_field[b["anchor"]["label"]] = b["mode"]
    assert by_field["보 길이 (m)"] == "number"
    assert by_field["배사구 높이"] == "number"
    assert by_field["하천명"] == "text"
    assert by_field["구조물번호"] == "text"              # 번호는 계산 대상이 아니다
    assert by_field["어도유무"] == "text"


def test_apply_writes_numbers_into_excel(tmp_path, monkeypatch):
    """4번 일괄 처리 결과 — 숫자 열은 엑셀에 수로 들어가 합계가 된다."""
    import app.main as app_main
    from fastapi.testclient import TestClient

    from core.pdf_pipeline import suggest_cells_maximal

    tpl = tmp_path / "tpl.pdf"
    _form(tpl, ROWS)
    # 값 칸만 남겨 템플릿으로(디자이너에서 라벨 칸 박스를 지운 상태) — 열 이름 = 항목명
    boxes = []
    for b in suggest_cells_maximal(str(tpl), 0):
        lbl = (b.get("anchor") or {}).get("label")
        if lbl and b["x0"] > 100:
            boxes.append({**b, "field": lbl})
    store = type("S", (), {"list_names": lambda self: ["조사표"],
                           "get": lambda self, n: {"name": "조사표", "boxes": boxes}})()
    monkeypatch.setattr(app_main, "_TEMPLATES", store)
    monkeypatch.setattr(app_main, "_tpl_pdf_path", lambda name: tpl)

    client = TestClient(app_main.app)
    with tpl.open("rb") as f:
        r = client.post("/api/pdf/apply",
                        data={"boxes": "[]", "sheet_name_field": "__group_title__",
                              "auto_classify": "1"},
                        files=[("files", ("in.pdf", f, "application/pdf"))])
    assert r.status_code == 200, r.text

    wb = openpyxl.load_workbook(app_main._PDF_APPLY["excel_path"])
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    vals = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(head)}
    assert vals["보 길이 (m)"] == 25.5                   # 단위 제거 + 수로 기록
    assert vals["배사구 높이"] == 1200                   # 천단위 콤마 제거 + 수
    assert isinstance(vals["보 길이 (m)"], float) and isinstance(vals["배사구 높이"], int)
    assert vals["하천명"] == "해남천"                     # 일반 열은 글자 그대로
    assert vals["구조물번호"] == "001"                    # 앞의 0이 살아 있다
