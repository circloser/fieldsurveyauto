"""데이터 입력 관리 — 템플릿 → 디지털 입력 양식 → 기록 가져오기 → 엑셀·현장조사표 PDF.

오토다타 웹(worker/index.js 의 Durable Object)은 여기서 같은 규칙의 가짜 서버로 대신한다
(양식 만들기 → 입력 키/관리 키, 기록 보내기 → seq, since 뒤의 기록만, 사진 따로, 닫기·삭제).
"""
from __future__ import annotations

import base64
import http.server
import io
import json
import re
import secrets
import threading
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import app.main as app_main
from core import forms
from core.template.store import TemplateStore

ORIGIN_RE = re.compile(r"^http://127\.0\.0\.1:\d+$")


class _Cloud(http.server.BaseHTTPRequestHandler):
    forms: dict = {}
    publish_key = ""
    port = 0

    def log_message(self, *a):  # noqa: D401
        pass

    def _json(self, code, data):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("content-type", "application/json; charset=utf-8")
        self.send_header("content-length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        n = int(self.headers.get("content-length") or 0)
        return json.loads(self.rfile.read(n).decode("utf-8")) if n else {}

    def _form(self, fid):
        return self.forms.get(fid)

    def do_POST(self):  # noqa: N802
        u = urllib.parse.urlparse(self.path)
        q = urllib.parse.parse_qs(u.query)
        if u.path == "/api/forms":
            # 클라우드플레어처럼 파이썬 기본 UA 는 봇으로 막는다(실서버에서 403 error code 1010 이 났던 회귀)
            if (self.headers.get("user-agent") or "").lower().startswith("python-urllib"):
                self.send_response(403)
                self.send_header("content-type", "text/plain")
                self.end_headers()
                self.wfile.write(b"error code: 1010\n")
                return None
            if self.publish_key and self.headers.get("x-publish-key") != self.publish_key:
                return self._json(403, {"error": "양식 만들기 키가 맞지 않습니다."})
            d = self._body()
            fid = "t" + secrets.token_hex(4)
            wk, ak = secrets.token_hex(8), secrets.token_hex(16)
            self.forms[fid] = {"definition": d["definition"], "write_key": wk, "admin_key": ak,
                               "closed": False, "entries": [], "photos": {}}
            return self._json(201, {"id": fid, "write_key": wk, "admin_key": ak,
                                    "share_url": f"http://127.0.0.1:{self.port}/f/{fid}?k={wk}"})
        m = re.match(r"^/api/forms/([a-z0-9]+)/(entries|close|open)$", u.path)
        if not m or not self._form(m.group(1)):
            return self._json(404, {"error": "없는 양식입니다."})
        f = self._form(m.group(1))
        if m.group(2) == "entries":
            if (q.get("k") or [""])[0] != f["write_key"]:
                return self._json(403, {"error": "링크의 키가 맞지 않습니다."})
            if f["closed"]:
                return self._json(409, {"error": "입력이 닫힌 양식입니다."})
            e = self._body()
            dup = next((x for x in f["entries"] if x["id"] == e["id"]), None)
            if dup:
                return self._json(200, {"ok": True, "seq": dup["seq"], "duplicate": True})
            seq = len(f["entries"]) + 1
            photos = []
            for field, data_url in (e.get("photos") or {}).items():
                mm = re.match(r"^data:(image/\w+);base64,(.+)$", data_url)
                raw = base64.b64decode(mm.group(2))
                f["photos"][(e["id"], field)] = (mm.group(1), raw)
                photos.append({"field": field, "mime": mm.group(1), "size": len(raw)})
            f["entries"].append({"seq": seq, "id": e["id"], "created": e.get("created"), "received": "2026-09-22T00:00:00Z",
                                 "client_id": e.get("client_id"), "values": e.get("values") or {},
                                 "meta": e.get("meta") or {}, "photos": photos})
            return self._json(201, {"ok": True, "seq": seq})
        if self.headers.get("x-admin-key") != f["admin_key"]:
            return self._json(403, {"error": "관리 키가 맞지 않습니다."})
        f["closed"] = m.group(2) == "close"
        return self._json(200, {"ok": True, "closed": f["closed"]})

    def do_GET(self):  # noqa: N802
        u = urllib.parse.urlparse(self.path)
        q = urllib.parse.parse_qs(u.query)
        m = re.match(r"^/api/forms/([a-z0-9]+)/entries$", u.path)
        if m:
            f = self._form(m.group(1))
            if not f:
                return self._json(404, {"error": "없는 양식입니다."})
            if self.headers.get("x-admin-key") != f["admin_key"]:
                return self._json(403, {"error": "관리 키가 맞지 않습니다."})
            since = int((q.get("since") or ["0"])[0])
            ents = [e for e in f["entries"] if e["seq"] > since]
            return self._json(200, {"entries": ents, "total": len(f["entries"]),
                                    "last_seq": len(f["entries"]), "closed": f["closed"], "definition": f["definition"]})
        m = re.match(r"^/api/forms/([a-z0-9]+)/photos/([^/]+)/(.+)$", u.path)
        if m:
            f = self._form(m.group(1))
            if not f or self.headers.get("x-admin-key") != f["admin_key"]:
                return self._json(403, {"error": "관리 키가 맞지 않습니다."})
            p = f["photos"].get((m.group(2), urllib.parse.unquote(m.group(3))))
            if not p:
                return self._json(404, {"error": "사진 없음"})
            self.send_response(200)
            self.send_header("content-type", p[0])
            self.send_header("content-length", str(len(p[1])))
            self.end_headers()
            self.wfile.write(p[1])
            return None
        return self._json(404, {"error": "없는 주소입니다."})

    def do_DELETE(self):  # noqa: N802
        m = re.match(r"^/api/forms/([a-z0-9]+)$", urllib.parse.urlparse(self.path).path)
        f = self._form(m.group(1)) if m else None
        if not f:
            return self._json(404, {"error": "없는 양식입니다."})
        if self.headers.get("x-admin-key") != f["admin_key"]:
            return self._json(403, {"error": "관리 키가 맞지 않습니다."})
        del self.forms[m.group(1)]
        return self._json(200, {"ok": True})


def _phone_submit(origin, fid, wk, entry):
    """휴대전화 입력 페이지가 하듯 기록 하나를 웹에 보낸다."""
    req = urllib.request.Request(f"{origin}/api/forms/{fid}/entries?k={wk}", method="POST",
                                 data=json.dumps(entry, ensure_ascii=False).encode("utf-8"),
                                 headers={"content-type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode("utf-8"))


BOXES = [
    {"field": "P1_제목", "page": 0, "x0": 150, "y0": 40, "x1": 450, "y1": 70, "mode": "title", "use_anchor": False,
     "anchor": {"label": "탐조 기록표", "relation": "right"}, "order": 1},
    {"field": "조사지역", "page": 0, "x0": 150, "y0": 100, "x1": 400, "y1": 122, "mode": "text", "use_anchor": True,
     "anchor": {"label": "조사 지역", "relation": "right"}, "order": 2},
    {"field": "개체수", "page": 0, "x0": 150, "y0": 130, "x1": 250, "y1": 152, "mode": "number", "use_anchor": True,
     "anchor": {"label": "개체수(마리)", "relation": "right"}, "order": 3},
    {"field": "야간조사", "page": 0, "x0": 300, "y0": 130, "x1": 322, "y1": 152, "mode": "check", "use_anchor": False,
     "anchor": {"label": "탐조 기록표", "relation": "right"}, "order": 4},
    {"field": "현장사진", "page": 0, "x0": 100, "y0": 170, "x1": 300, "y1": 320, "mode": "image", "order": 5},
    {"field": "관찰표", "page": 0, "x0": 100, "y0": 400, "x1": 500, "y1": 500, "mode": "table",
     "columns": ["시각", "행동", "비고"], "order": 6},
]


def _make_template_pdf(path: Path) -> None:
    """라벨과 표 칸 선이 있는 양식 PDF 한 쪽(박스 좌표와 맞춤)."""
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    font = fitz.Font("cjk")
    tw = fitz.TextWriter(page.rect)
    tw.append((150, 62), "탐조 기록표", font=font, fontsize=16)
    tw.append((60, 116), "조사 지역", font=font, fontsize=10)
    tw.append((60, 146), "개체수(마리)", font=font, fontsize=10)
    tw.append((260, 146), "야간", font=font, fontsize=10)
    tw.append((60, 190), "현장 사진", font=font, fontsize=10)
    for j, c in enumerate(["시각", "행동", "비고"]):
        tw.append((105 + j * 133, 417), c, font=font, fontsize=10)
    tw.append((155, 116), "예시값", font=font, fontsize=9)          # '작성 예시' 양식 — 값·√ 가 이미 인쇄됨
    tw.append((306, 146), "√", font=font, fontsize=10)
    tw.write_text(page)
    for r in (fitz.Rect(150, 100, 400, 122), fitz.Rect(150, 130, 250, 152), fitz.Rect(300, 130, 322, 152),
              fitz.Rect(100, 170, 300, 320)):
        page.draw_rect(r, color=(0, 0, 0), width=0.8)
    xs = [100, 233, 366, 500]
    for i in range(4):                       # 머리글 1줄 + 데이터 3줄
        for j in range(3):
            page.draw_rect(fitz.Rect(xs[j], 400 + i * 25, xs[j + 1], 425 + i * 25), color=(0, 0, 0), width=0.8)
    doc.save(str(path))
    doc.close()


def _jpeg() -> bytes:
    from PIL import Image
    im = Image.new("RGB", (320, 240), (40, 130, 210))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=75)
    return buf.getvalue()


@pytest.fixture
def cloud(tmp_path, monkeypatch):
    _Cloud.forms = {}
    _Cloud.publish_key = ""
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Cloud)
    _Cloud.port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    origin = f"http://127.0.0.1:{_Cloud.port}"
    monkeypatch.setenv("AUTODATA_FORMS_ORIGIN", origin)
    monkeypatch.delenv("AUTODATA_FORMS_PUBLISH_KEY", raising=False)
    monkeypatch.setattr(app_main, "_FORMS", forms.FormRegistry(tmp_path / "forms.json", tmp_path / "forms"))
    monkeypatch.setattr(app_main, "_TEMPLATES", TemplateStore(tmp_path / "templates.json"))
    monkeypatch.setattr(app_main.config, "TEMPLATE_PDF_DIR", tmp_path / "tpdf")
    monkeypatch.setattr(app_main.config, "OUTPUT_DIR", tmp_path / "out")
    (tmp_path / "tpdf").mkdir()
    (tmp_path / "out").mkdir()
    app_main._TEMPLATES.save("시험양식", BOXES)
    _make_template_pdf(app_main._tpl_pdf_path("시험양식"))
    yield {"origin": origin, "registry": app_main._FORMS, "tmp": tmp_path}
    httpd.shutdown()


def test_build_definition_types_labels_and_columns():
    d = forms.build_definition("시험양식", BOXES, title="", gps=True)
    assert d["title"] == "시험양식" and d["gps"] is True and d["template"] == "시험양식"
    by = {f["key"]: f for f in d["fields"]}
    assert by["P1_제목"]["type"] == "title" and by["P1_제목"]["value"] == "탐조 기록표"
    assert by["조사지역"]["label"] == "조사 지역"          # 라벨 기준 박스 → 인쇄된 라벨
    assert by["야간조사"]["label"] == "야간조사"           # 라벨을 안 쓰는 박스 → 박스 이름
    assert by["개체수"]["type"] == "number" and by["야간조사"]["type"] == "check"
    assert by["현장사진"]["type"] == "image" and by["관찰표"]["columns"] == ["시각", "행동", "비고"]
    only_title = forms.build_definition("x", [dict(BOXES[0]), dict(BOXES[5], columns=[])], gps=False)
    assert only_title["fields"][1]["columns"] == ["값"]
    with pytest.raises(forms.FormsError):
        forms.build_definition("x", [BOXES[0]])


def test_duplicate_and_placeholder_boxes():
    # 같은 이름 박스 → 일괄 처리와 같은 '이름 (2)' 규칙, 이름 없는 제목 박스('새 항목')는 입력 페이지에 안 보임
    boxes = [dict(BOXES[0], field="새 항목", anchor=None), dict(BOXES[1], field="비고", order=2),
             dict(BOXES[1], field="비고", order=3, page=1), dict(BOXES[0], field="새 항목", anchor=None, order=4, page=1)]
    dd = forms.dedup_boxes(boxes)
    assert [b["field"] for b in dd] == ["새 항목", "비고", "비고 (2)", "새 항목 (2)"]
    d = forms.build_definition("x", dd)
    assert [f["key"] for f in d["fields"]] == ["비고", "비고 (2)"]


def test_settings_from_env_and_config_file(tmp_path, monkeypatch):
    monkeypatch.delenv("AUTODATA_FORMS_ORIGIN", raising=False)
    monkeypatch.delenv("AUTODATA_FORMS_PUBLISH_KEY", raising=False)
    assert forms.settings(tmp_path) == {"origin": "https://autodata.singlena.workers.dev", "publish_key": ""}
    (tmp_path / forms.CONFIG_FILE).write_text("# 주석\nORIGIN = https://forms.example.org/\npublish_key=abc\n", encoding="utf-8")
    assert forms.settings(tmp_path) == {"origin": "https://forms.example.org", "publish_key": "abc"}
    monkeypatch.setenv("AUTODATA_FORMS_ORIGIN", "http://127.0.0.1:8787")
    assert forms.settings(tmp_path)["origin"] == "http://127.0.0.1:8787"


def test_cloud_client_sends_own_user_agent(cloud):
    seen = {}
    orig = _Cloud.do_POST

    def spy(self):
        seen["ua"] = self.headers.get("user-agent")
        return orig(self)
    _Cloud.do_POST = spy
    try:
        forms.CloudClient(cloud["origin"]).create({"title": "x", "fields": [{"key": "a", "type": "text"}]})
    finally:
        _Cloud.do_POST = orig
    assert seen["ua"].startswith("AutoData-helper")


def test_qr_svg_and_local_time():
    svg = forms.qr_svg("https://autodata.singlena.workers.dev/f/abcdefghij?k=0123456789abcdef")
    assert svg.startswith("<svg") and "path" in svg
    assert re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$", forms.local_time("2026-09-22T10:00:00.000Z"))
    assert forms.local_time("") == "" and forms.local_time("이상한값") == "이상한값"


def test_api_rejects_bad_ids(cloud):
    c = TestClient(app_main.app)
    assert c.get("/api/forms/ZZ/entries").status_code == 404
    assert c.post("/api/forms/nope99/sync").status_code == 404
    assert c.post("/api/forms/publish", json={"template": "없는템플릿"}).status_code == 404


def test_publish_sync_export_roundtrip(cloud):
    c = TestClient(app_main.app)
    origin, reg = cloud["origin"], cloud["registry"]

    # 1. 양식 만들기 → 공유 링크(웹 주소) · 관리 키는 화면에 안 줌 · 스냅숏·양식 PDF 보관
    r = c.post("/api/forms/publish", json={"template": "시험양식", "title": "탐조 기록(시험)", "gps": True})
    assert r.status_code == 200, r.text
    form = r.json()["form"]
    fid = form["id"]
    assert form["share_url"].startswith(origin + "/f/") and "admin_key" not in form and form["has_pdf"]
    assert form["title"] == "탐조 기록(시험)" and form["definition"]["gps"] is True
    assert (reg.root / fid / "template.pdf").exists() and (reg.root / fid / "template.json").exists()
    listed = c.get("/api/forms").json()
    assert [f["id"] for f in listed["forms"]] == [fid] and listed["origin"] == origin
    qr = c.get(f"/api/forms/{fid}/qr.svg")
    assert qr.status_code == 200 and qr.headers["content-type"].startswith("image/svg+xml")

    # 2. 현장(휴대전화)에서 기록 2건 — 하나는 사진·표·위치 포함
    wk = form["write_key"]
    jpg = _jpeg()
    e1 = {"id": "entry0000000001", "client_id": "phoneA", "created": "2026-09-22T10:00:00.000Z",
          "values": {"조사지역": "남동유수지", "개체수": 12, "야간조사": True,
                     "관찰표": [{"시각": "09:10", "행동": "채식", "비고": ""},
                                {"시각": "09:40", "행동": "휴식", "비고": "둥지 근처"}]},
          "photos": {"현장사진": "data:image/jpeg;base64," + base64.b64encode(jpg).decode()},
          "meta": {"gps": {"lat": 37.4, "lon": 126.7, "acc": 8}}}
    e2 = {"id": "entry0000000002", "client_id": "phoneB", "created": "2026-09-22T10:05:00.000Z",
          "values": {"조사지역": "송도갯벌", "개체수": 3, "야간조사": False, "관찰표": []}, "photos": {}, "meta": {}}
    assert _phone_submit(origin, fid, wk, e1)[0] == 201
    assert _phone_submit(origin, fid, wk, e2)[0] == 201

    # 3. 가져오기 → 기록·사진이 이 PC에
    r = c.post(f"/api/forms/{fid}/sync")
    assert r.status_code == 200, r.text
    assert r.json()["new"] == 2 and r.json()["total"] == 2 and r.json()["form"]["count"] == 2
    d = c.get(f"/api/forms/{fid}/entries").json()
    ents = d["entries"]
    assert [e["seq"] for e in ents] == [1, 2]
    assert ents[0]["values"]["조사지역"] == "남동유수지" and ents[0]["values"]["관찰표"][1]["비고"] == "둥지 근처"
    assert re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$", ents[0]["created_local"])
    photo_url = ents[0]["photos"]["현장사진"]
    assert photo_url == f"/api/forms/{fid}/photo/entry0000000001/현장사진"
    p = c.get(photo_url)
    assert p.status_code == 200 and p.content == jpg
    assert (reg.root / fid / "photos").exists()
    assert c.post(f"/api/forms/{fid}/sync").json()["new"] == 0          # 다시 가져와도 새 기록 없음

    # 4. 엑셀 — 표 기록은 줄마다 한 행, 숫자·위도는 수, 체크는 항목 이름, 사진은 그림
    r = c.get(f"/api/forms/{fid}/excel")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"], r.text
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    assert headers == ["기록", "기록일시", "위도", "경도", "위치오차m", "P1_제목", "조사지역", "개체수", "야간조사",
                       "현장사진", "시각", "행동", "비고"]
    rows = [[ws.cell(row=r_, column=i).value for i in range(1, ws.max_column + 1)] for r_ in range(2, ws.max_row + 1)]
    assert len(rows) == 3                                        # 기록 1 = 표 2줄, 기록 2 = 1줄
    col = {h: i for i, h in enumerate(headers)}
    assert rows[0][col["기록"]] == "기록 1" and rows[1][col["기록"]] == "기록 1" and rows[2][col["기록"]] == "기록 2"
    assert rows[0][col["조사지역"]] == rows[1][col["조사지역"]] == "남동유수지"
    assert rows[0][col["개체수"]] == 12 and isinstance(rows[0][col["개체수"]], int)
    assert rows[0][col["위도"]] == 37.4 and rows[2][col["위도"]] is None
    assert rows[0][col["야간조사"]] == "야간조사" and rows[2][col["야간조사"]] is None
    assert rows[0][col["P1_제목"]] == "탐조 기록표"
    assert [rows[0][col["시각"]], rows[1][col["시각"]]] == ["09:10", "09:40"] and rows[1][col["비고"]] == "둥지 근처"
    assert rows[0][col["현장사진"]] is None and len(ws._images) == 1  # 그림으로 들어감

    # 5. 현장조사표 PDF — 기록마다 양식 한 벌, 값·체크·표·바닥글이 글자로 있음
    r = c.get(f"/api/forms/{fid}/pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf", r.text
    import fitz
    pdf = fitz.open(stream=r.content, filetype="pdf")
    assert pdf.page_count == 2
    t0 = pdf[0].get_text()
    for s in ("남동유수지", "12", "√", "09:40", "휴식", "둥지 근처", "기록 1", "37.4"):
        assert s in t0, s
    assert len(pdf[0].get_images()) >= 1                         # 사진 삽입
    t1 = pdf[1].get_text()
    assert "송도갯벌" in t1 and "기록 2" in t1        # (인쇄된 예시 √ 는 글자로는 남고 눈에는 덮임 — 아래 픽셀 확인)
    one = c.get(f"/api/forms/{fid}/pdf", params={"entry": "entry0000000002"})
    assert fitz.open(stream=one.content, filetype="pdf").page_count == 1
    # 작성 예시 양식의 인쇄된 값·√ 는 흰색으로 덮인다(글자 추출에는 남지만 눈에는 안 보임) — 픽셀로 확인
    def dark(pg, x0, y0, x1, y1):
        pix = pg.get_pixmap(dpi=144, clip=fitz.Rect(x0, y0, x1, y1))
        return sum(1 for i in range(0, len(pix.samples), pix.n) if pix.samples[i] < 128)
    assert dark(pdf[1], 302, 132, 320, 150) == 0                 # 기록 2(체크 안 함): 인쇄된 √ 가 덮여 흰 칸
    assert dark(pdf[0], 302, 132, 320, 150) > 20                 # 기록 1(체크): √ 가 그려짐
    assert dark(pdf[0], 152, 102, 398, 120) > 50                 # 값 '남동유수지'는 보인다

    # 6. 입력 닫기 → 현장에서 더 못 보냄 → 다시 열기
    r = c.post(f"/api/forms/{fid}/close")
    assert r.status_code == 200 and r.json()["form"]["closed"] is True
    assert _phone_submit(origin, fid, wk, dict(e2, id="entry0000000003"))[0] == 409
    assert c.post(f"/api/forms/{fid}/open").json()["form"]["closed"] is False
    assert _phone_submit(origin, fid, wk, dict(e2, id="entry0000000003"))[0] == 201

    # 7. 삭제 → 웹·이 PC 모두 비움
    r = c.post(f"/api/forms/{fid}/delete")
    assert r.status_code == 200 and r.json()["forms"] == []
    assert fid not in _Cloud.forms and not (reg.root / fid).exists()
    assert c.get(f"/api/forms/{fid}/entries").status_code == 404


def test_publish_without_pdf_gives_excel_only(cloud):
    c = TestClient(app_main.app)
    app_main._TEMPLATES.save("PDF없음", BOXES[:4])
    r = c.post("/api/forms/publish", json={"template": "PDF없음"})
    assert r.status_code == 200 and r.json()["form"]["has_pdf"] is False
    fid = r.json()["form"]["id"]
    r = c.get(f"/api/forms/{fid}/pdf")
    assert r.status_code == 502 and "양식 PDF" in r.json()["error"]
    assert c.get(f"/api/forms/{fid}/excel").status_code == 200        # 기록이 없어도 머리글만 있는 엑셀


def test_cloud_errors_are_readable(cloud, monkeypatch):
    c = TestClient(app_main.app)
    _Cloud.publish_key = "secret"
    r = c.post("/api/forms/publish", json={"template": "시험양식"})
    assert r.status_code == 400 and "키" in r.json()["error"]
    monkeypatch.setenv("AUTODATA_FORMS_PUBLISH_KEY", "secret")
    assert c.post("/api/forms/publish", json={"template": "시험양식"}).status_code == 200
    monkeypatch.setenv("AUTODATA_FORMS_ORIGIN", "http://127.0.0.1:1")    # 아무도 없는 포트
    r = c.post("/api/forms/publish", json={"template": "시험양식"})
    assert r.status_code == 502 and "연결할 수 없습니다" in r.json()["error"]
